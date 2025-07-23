import pandas as pd
from docx import Document
from datetime import datetime

# Charger le fichier Excel
fichier = "data/daily_data.xlsx"
df = pd.read_excel(fichier)

doc = Document()
doc.add_heading(f'Daily Meeting - {datetime.now().strftime("%d %B %Y")}', 0)

table = doc.add_table(rows=1, cols=5)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Membre'
hdr_cells[1].text = 'Tâches de la veille'
hdr_cells[2].text = 'Tâches du jour'
hdr_cells[3].text = 'Difficultés'
hdr_cells[4].text = 'Tâches de demain'

for _, row in df.iterrows():
    membre = str(row.get('Nom', row.get('Auteur', '')))
    veille = str(row.get('Travaux de la veille', row.get('Tâche veille', '')))
    # Tâches du jour : concatène chaque tâche avec avancement si dispo
    taches_jour = row.get('Travaux du jour', row.get('Tâche jour', ''))
    if isinstance(taches_jour, str):
        taches_jour_str = taches_jour
    else:
        taches_jour_str = '\n'.join(
            f"{t['tache']} (Avancement : {t['avancement']})" if t.get('avancement') else t['tache']
            for t in taches_jour if isinstance(t, dict)
        )
    diff = str(row.get('Difficultés', ''))
    taches_demain = str(row.get('Taches de demain', ''))

    cells = table.add_row().cells
    cells[0].text = membre
    cells[1].text = veille
    cells[2].text = taches_jour_str
    cells[3].text = diff
    cells[4].text = taches_demain

doc.save("Daily_Meeting_Report.docx")
print("Rapport Word généré.")
